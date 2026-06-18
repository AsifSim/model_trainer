# import json
# import logging
# from pathlib import Path

# from docx import Document

# logging.basicConfig(
#     level=logging.INFO,
#     format="%(asctime)s %(levelname)s %(message)s"
# )

# logger = logging.getLogger(__name__)

# ROOT = Path("TBRD_CONVERSION")
# OUTPUT = "dataset.jsonl"

# SYSTEM_PROMPT = """
# You are an expert Java developer. Your task is to convert the provided TBRD requirement document into a clean, syntactically correct, and executable Java method.

# Rules:
# - Generate ONLY valid Java code inside the markdown code block.
# - Do not include conversational explanations, preambles, or notes.
# - Ensure all fields, parameters, and entities from the requirement text are preserved in the method signature or logic.
# - Do not hallucinate fields or logic not present in the requirement text.
# - Use clean formatting, logical naming, and standard Java code conventions.
# """


# def extract_docx_content(
#         path: Path
# ) -> str:

#     document = Document(path)

#     content = []

#     extract_paragraphs(
#         document,
#         content
#     )

#     extract_tables(
#         document,
#         content
#     )

#     return "\n".join(content)


# def extract_paragraphs(
#         document,
#         content: list[str]
# ) -> None:

#     for paragraph in document.paragraphs:

#         text = paragraph.text.strip()

#         if text:

#             content.append(text)


# def extract_tables(
#         document,
#         content: list[str]
# ) -> None:

#     for table in document.tables:

#         content.append("[TABLE]")

#         for row in table.rows:

#             values = [
#                 cell.text.strip()
#                 for cell in row.cells
#             ]

#             content.append(
#                 " | ".join(values)
#             )

#         content.append("[/TABLE]")


# def read_target_text(
#         path: Path
# ) -> str:

#     return path.read_text(
#         encoding="utf-8",
#         errors="ignore"
#     ).strip()


# def create_record(
#         source_text: str,
#         target_text: str
# ) -> dict:

#     return {
#         "messages": [
#             {
#                 "role": "system",
#                 "content": SYSTEM_PROMPT
#             },
#             {
#                 "role": "user",
#                 "content":
#                     "Input TBRD Document:\n\n"
#                     + source_text
#             },
#             {
#                 "role": "assistant",
#                 "content": target_text
#             }
#         ]
#     }


# def process_document_pair(
#         docx_file: Path
# ) -> dict | None:

#     txt_file = (
#             ROOT /
#             f"{docx_file.stem}.txt"
#     )

#     if not txt_file.exists():

#         logger.warning(
#             "Missing target file %s",
#             txt_file.name
#         )

#         return None

#     logger.info(
#         "Processing %s",
#         docx_file.stem
#     )

#     source_text = extract_docx_content(
#         docx_file
#     )

#     target_text = read_target_text(
#         txt_file
#     )

#     return create_record(
#         source_text,
#         target_text
#     )


# def build_dataset():

#     with open(
#             OUTPUT,
#             "w",
#             encoding="utf-8"
#     ) as writer:

#         for docx_file in ROOT.glob("*.docx"):

#             record = process_document_pair(
#                 docx_file
#             )

#             if record is None:
#                 continue

#             writer.write(
#                 json.dumps(
#                     record,
#                     ensure_ascii=False
#                 )
#                 + "\n"
#             )

#     logger.info(
#         "Dataset written to %s",
#         OUTPUT
#     )


# if __name__ == "__main__":
#     build_dataset()
import json
import logging
from pathlib import Path
from docx import Document

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s"
)

logger = logging.getLogger(__name__)

# --- Configuration ---
ROOT = Path("JSON_CREATION")        # Folder containing your .docx and .json pairs
OUTPUT = "dataset_json.jsonl"      # Final output dataset file

SYSTEM_PROMPT = """
You are an expert data architect and deterministic JSON compiler.

Your task is to convert plain English business or functional requirements into a structured JSON configuration following the target microservice schema.

OUTPUT RULES

* Output ONLY valid JSON.
* Output exactly one JSON object.
* Do not output markdown.
* Do not output code fences.
* Do not output explanations.
* Do not output notes.
* Do not output examples.
* Do not output schema descriptions.
* Do not output any text before or after the JSON.
* The response must begin with '{' and end with '}'.

ROOT STRUCTURE

Always generate:

{
"microservice": {},
"requestFields": [],
"dbCalls": [],
"calculations": [],
"finalOutput": {}
}

MICROSERVICE RULES

Generate:

{
"microservice": {
"name": "<service name>",
"type": "PLATFORM_TRIGGERED",
"version": "1.0"
}
}

* Derive the service name from the business requirement.
* Default type to PLATFORM_TRIGGERED unless specified.
* Default version to 1.0 unless specified.

REQUEST FIELD RULES

* Extract all request-level input fields.
* Include only fields supplied by the incoming request.
* Exclude derived fields.
* Exclude database outputs.
* Remove duplicates.

DB CALL RULES

Each unique data retrieval operation becomes one dbCall.

dbCall structure:

{
"id": number,
"name": string,
"consumes": [],
"entity": string,
"streaming": boolean,
"dynamicJoinModel": object|null,
"joins": [],
"produces": []
}

* Assign sequential ids.
* Never create duplicate dbCalls.
* Reuse dbCalls when the same lookup is referenced multiple times.
* consumes contains dependencies required before execution.
* produces contains fields returned by the lookup.
* entity is the source entity being queried.

LOOKUP MAPPING

When requirements mention:

* retrieve
* fetch
* lookup
* load
* get

create a dbCall.

Simple joins use:

{
"leftField": "<entity field>",
"rightSource": "<source.field>"
}

DYNAMIC JOIN MAPPING

When requirements mention:

* join
* combine
* correlate
* merge
* enrich
* aggregate across entities

generate a dynamicJoinModel.

Structure:

{
"anchorEntity": "",
"joins": [],
"filters": [],
"returnFields": []
}

Join format:

{
"from": "entity1.field",
"to": "entity2.field"
}

Filter format:

{
"field": "",
"operator": "=",
"valueSource": ""
}

CALCULATION RULES

Create a calculation whenever requirements mention:

* calculate
* compute
* determine
* evaluate
* validate
* assess
* score
* aggregate

Calculation structure:

{
"name": "",
"consumes": [],
"steps": []
}

PREPROCESS STEP

{
"type": "PREPROCESS",
"description": ""
}

CONDITIONAL LOGIC

Convert business rules into IF structures.

Structure:

{
"type": "IF",
"condition": "",
"then": [],
"elseIf": [],
"else": []
}

* Preserve exact conditions.
* Preserve thresholds.
* Preserve AND / OR logic.
* Preserve comparison operators.

ACTION MAPPING

Create record -> INSERT

{
"type": "INSERT",
"entity": "",
"fieldMappings": {}
}

Update or modify record -> UPDATE

{
"type": "UPDATE",
"entity": "",
"targetRecord": {},
"fieldMappings": {}
}

Delete record -> DELETE

{
"type": "DELETE",
"entity": "",
"targetRecord": {}
}

Invoke service -> TRIGGER_SERVICE

{
"type": "TRIGGER_SERVICE",
"service": "",
"inputs": []
}

FIELD MAPPINGS

Map destination fields to source values.

Example:

{
"status": "CalculationResult.status",
"discount": "LoyaltyProgram.discountPercentage"
}

FINAL OUTPUT RULES

Always generate finalOutput.

Structure:

{
"type": "UPDATE",
"entity": "",
"targetRecord": {},
"fieldMappings": {}
}

* Represents the final state mutation.
* References calculation outputs when applicable.
* Contains final business outcome fields.

VALIDATION RULES

Before generating JSON:

* Ensure all root sections exist.
* Ensure dbCall ids are unique.
* Ensure calculation names are unique.
* Ensure all consumed dependencies exist.
* Ensure all referenced fields exist.
* Ensure UPDATE actions contain targetRecord.
* Ensure TRIGGER_SERVICE actions contain service.
* Ensure dynamic joins contain at least one join.
* Ensure no duplicate lookups exist.
* Ensure no duplicate calculations exist.
* Ensure output is valid JSON.

"""

def extract_docx_content(path: Path) -> str:
    """Extracts all text paragraphs and tables from a .docx file."""
    document = Document(path)
    content = []

    # 1. Pull paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            content.append(text)

    # 2. Pull tables (if any exist in your requirements)
    for table in document.tables:
        content.append("[TABLE]")
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            content.append(" | ".join(values))
        content.append("[/TABLE]")

    return "\n".join(content)

def read_target_json(path: Path) -> str:
    """Reads the expected target JSON payload file."""
    return path.read_text(encoding="utf-8", errors="ignore").strip()

def create_record(source_text: str, target_json: str) -> dict:
    """Structures data into standard ChatML format for Unsloth training."""
    return {
        "messages": [
            {
                "role": "system",
                "content": SYSTEM_PROMPT
            },
            {
                "role": "user",
                "content": "Convert this requirement into a clean JSON structure:\n\n" + source_text
            },
            {
                "role": "assistant",
                "content": target_json
            }
        ]
    }

def process_document_pair(docx_file: Path) -> dict | None:
    """Pairs a .docx file with its matching .json file based on the file name."""
    json_file = ROOT / f"{docx_file.stem}.json"

    if not json_file.exists():
        logger.warning(
            "Missing matching target file: %s (Expected based on %s)",
            json_file.name, docx_file.name
        )
        return None

    logger.info("Processing pair: %s", docx_file.stem)
    
    source_text = extract_docx_content(docx_file)
    target_json = read_target_json(json_file)

    return create_record(source_text, target_json)

def build_dataset():
    if not ROOT.exists():
        logger.error("Directory '%s' not found. Please create it and add your file pairs.", ROOT)
        return

    counter = 0
    with open(OUTPUT, "w", encoding="utf-8") as writer:
        # Loop over all Word files instead of text files
        for docx_file in ROOT.glob("*.docx"):
            record = process_document_pair(docx_file)

            if record is None:
                continue

            writer.write(json.dumps(record, ensure_ascii=False) + "\n")
            counter += 1

    logger.info("🎉 Dataset successfully written to %s with %d records!", OUTPUT, counter)

if __name__ == "__main__":
    build_dataset()